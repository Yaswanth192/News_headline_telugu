from pathlib import Path
from pptx import Presentation
from pptx.util import Inches, Pt

# Paths
out_dir = Path(r'D:\large documents\Large Projects\News_Headline\submission')
ppt_path = out_dir / 'Telugu_News_Headline_Generation_Presentation.pptx'
notes_path = out_dir / 'PRESENTATION_SPEAKER_NOTES.md'

# Create presentation
prs = Presentation()

# Helper to add slide with title and bullet points
def add_slide(title, bullets=None):
    slide_layout = prs.slide_layouts[1]  # Title and Content
    slide = prs.slides.add_slide(slide_layout)
    slide.shapes.title.text = title
    if bullets:
        tf = slide.placeholders[1].text_frame
        for b in bullets:
            p = tf.add_paragraph()
            p.text = b
            p.level = 0
    return slide

# Slide contents (simplified placeholders)
slides_info = [
    ("Title", ["Telugu News Headline Generation", "Final Project Report", "Generated on 2026-10-13"]),
    ("Problem & Motivation", ["Automatic headline generation for Telugu news", "Support low‑resource settings", "Improve editorial efficiency"]),
    ("Objectives", ["Train and compare mT5‑small & IndicBART", "Build relevance classifier", "Low‑resource experiments", "Report findings"]),
    ("Dataset", ["Source: saidines12/telugu_news_dataset", "~83k train, 5.9k val, 12.5k test", "Duplicate & leakage analysis included"]),
    ("Model Architecture", ["Seq2Seq mT5‑small (300M params)", "IndicBART (244M params)", "BERT relevance classifier"]),
    ("Training Setup", ["bfloat16, batch=1, grad_accum=16", "50 steps (full), 75 steps (low‑resource)", "Runs on RTX 3050 (4 GB VRAM)"]),
    ("Results – Model Comparison", ["mT5‑small ROUGE‑1=0.1025, latency=0.27 s", "IndicBART ROUGE‑1=0.0109, latency=2.62 s"]),
    ("Low‑Resource Findings", ["5 % data retains ~85 % ROUGE‑1 of full data", "Training time comparable across tiers"]),
    ("Relevance Classification", ["mBERT Macro‑F1=0.5114 (Accuracy=0.5238)", "TF‑IDF+LR Macro‑F1=0.4775"]),
    ("Gemma 2 Feasibility", ["Blocked: HF gated repo, missing bitsandbytes/peft, VRAM insufficient"]),
    ("Streamlit Demo", ["Headline generation page", "Relevance analysis page", "Dashboard with metrics"]),
    ("Limitations", ["Very short training (50‑75 steps)", "Small evaluation sample sizes", "No human evaluation", "Gemma 2 not trained"]),
    ("Conclusion", ["mT5‑small is the pragmatic choice for low‑resource Telugu headline generation", "Future work: longer training, human studies, larger LLMs"]),
]

for title, bullets in slides_info:
    add_slide(title, bullets)

# Save PPTX
prs.save(str(ppt_path))
print('PPTX saved to', ppt_path)

# Create speaker notes markdown
with notes_path.open('w', encoding='utf-8') as f:
    for idx, (title, bullets) in enumerate(slides_info, start=1):
        f.write(f"## Slide {idx}: {title}\n\n")
        for b in bullets:
            f.write(f"- {b}\n")
        f.write("\n---\n\n")
print('Speaker notes saved to', notes_path)
